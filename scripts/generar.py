# -*- coding: utf-8 -*-
"""Regenera index.html a partir del Excel de control.

Lee el archivo que esté en datos/, arma la información de las plantillas y la
inyecta en plantilla/tablero.html entre los marcadores SEED, HIST y SELLO.
Las etapas que se siguen se toman de ETAPAS_CFG, dentro de la propia plantilla,
para que agregar una etapa siga siendo un solo cambio en un solo lugar.
"""
import datetime, glob, json, os, re, subprocess, sys
from zoneinfo import ZoneInfo

import openpyxl

RAIZ      = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANTILLA = os.path.join(RAIZ, "plantilla", "tablero.html")
SALIDA    = os.path.join(RAIZ, "index.html")
HISTORICO = os.path.join(RAIZ, "datos", "historico.json")
TZ        = ZoneInfo("America/Mexico_City")
CERRADAS  = ("completado", "no aplica", "n/a")


def archivo_excel():
    """El Excel de control. Si hay varios, gana el del commit más reciente."""
    rutas = [r for r in sorted(glob.glob(os.path.join(RAIZ, "datos", "*.xlsx")))
             if not os.path.basename(r).startswith("~$")]
    if not rutas:
        sys.exit("No hay ningún archivo .xlsx en la carpeta datos/.")
    if len(rutas) > 1:
        def commit(r):
            p = subprocess.run(["git", "log", "-1", "--format=%ct", "--", r],
                               cwd=RAIZ, capture_output=True, text=True)
            return int(p.stdout.strip() or 0)
        rutas.sort(key=commit)
    return rutas[-1]


def texto(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def columnas_de_etapa(plantilla):
    bloque = re.search(r"const ETAPAS_CFG = \[([\s\S]*?)\n\];", plantilla)
    cols = re.findall(r'col:"([^"]+)"', bloque.group(1)) if bloque else []
    return cols or ["E1 Llenado", "E2 Revisión", "E3 Repositorio"]


def leer_plantillas(ruta, cols_etapa):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    nombre = next((h for h in wb.sheetnames if h.strip().lower().startswith("control")),
                  wb.sheetnames[0])
    filas = list(wb[nombre].iter_rows(values_only=True))

    enc = next((i for i, f in enumerate(filas)
                if f and any(texto(c).lower() == "plantilla" for c in f)), None)
    if enc is None:
        sys.exit("No se encontró el encabezado: la hoja de control debe conservar sus títulos.")
    H = [texto(c).lower() for c in filas[enc]]

    def col(*alternativas):
        for a in alternativas:
            a = a.strip().lower()
            if a in H:
                return H.index(a)
        return -1

    c = {"id": col("ID"), "modulo": col("Módulo", "Modulo"), "codigo": col("Código", "Codigo"),
         "plantilla": col("Plantilla"), "area": col("Área Responsable", "Area Responsable"),
         "keyuser": col("Key User"), "revisor": col("Revisor / Consultor"),
         "prioridad": col("Prioridad"), "ultmod": col("Últ. Modificación", "Ult. Modificacion"),
         "obs": col("Observaciones / Bloqueos", "Observaciones")}

    etapas = []
    for i, nombre_col in enumerate(cols_etapa, start=1):
        est = col(nombre_col)
        if est < 0:
            sys.exit("Falta la columna de estatus «%s» en la hoja de control." % nombre_col)
        etapas.append((est, col("E%d F. Compromiso" % i), col("E%d F. Real" % i)))

    def valor(fila, indice):
        return texto(fila[indice]) if 0 <= indice < len(fila) else ""

    plantillas = []
    for n, fila in enumerate(filas[enc + 1:], start=enc + 2):
        if not fila or not valor(fila, c["plantilla"]):
            continue
        plantillas.append({
            "id":        valor(fila, c["id"]) or "R%d" % n,
            "modulo":    valor(fila, c["modulo"]) or "Sin módulo",
            "codigo":    valor(fila, c["codigo"]),
            "plantilla": valor(fila, c["plantilla"]),
            "area":      valor(fila, c["area"]),
            "keyuser":   valor(fila, c["keyuser"]),
            "revisor":   valor(fila, c["revisor"]),
            "prioridad": valor(fila, c["prioridad"]),
            "ultmod":    valor(fila, c["ultmod"]),
            "obs":       valor(fila, c["obs"]),
            "e": [{"est": valor(fila, est) or "No iniciado",
                   "comp": valor(fila, comp), "real": valor(fila, real)}
                  for est, comp, real in etapas],
        })
    if not plantillas:
        sys.exit("La hoja de control no tiene filas de plantillas.")
    return plantillas


def conteos(plantillas):
    """Mismo criterio que el tablero: entregada = última etapa vigente concluida."""
    c = {"total": len(plantillas), "ent": 0, "sin": 0, "proceso": 0}
    for r in plantillas:
        estatus = [e["est"].strip().lower() for e in r["e"]]
        if estatus[-1] in CERRADAS:
            c["ent"] += 1
        elif all(e == "no iniciado" for e in estatus):
            c["sin"] += 1
        else:
            c["proceso"] += 1
    return c


def historico(actual, dia):
    previos = []
    if os.path.exists(HISTORICO):
        with open(HISTORICO, encoding="utf-8") as f:
            previos = [h for h in json.load(f) if h.get("fecha") != dia]
    previos.append({"fecha": dia, **actual})
    previos = previos[-12:]
    with open(HISTORICO, "w", encoding="utf-8") as f:
        json.dump(previos, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return previos


def seguro(dato):
    """JSON que puede vivir dentro de un <script> sin romperlo."""
    return json.dumps(dato, ensure_ascii=False, separators=(",", ":")) \
               .replace("*/", "*\\/").replace("</", "<\\/")


def main():
    with open(PLANTILLA, encoding="utf-8") as f:
        plantilla = f.read()

    ruta = archivo_excel()
    filas = leer_plantillas(ruta, columnas_de_etapa(plantilla))

    ahora = datetime.datetime.now(TZ)
    dia   = ahora.strftime("%d/%m/%Y")
    # el nombre que se muestra es el del archivo en Drive, no el del archivo local
    nombre_drive = os.path.basename(ruta)
    if os.path.exists(os.path.join(RAIZ, "datos", "origen.json")):
        with open(os.path.join(RAIZ, "datos", "origen.json"), encoding="utf-8") as f:
            nombre_drive = json.load(f).get("nombre") or nombre_drive
    sello = {"archivo": nombre_drive, "fecha": ahora.strftime("%d/%m/%Y %H:%M")}
    hist  = historico(conteos(filas), dia)

    html = plantilla
    for marca, dato in (("SEED", filas), ("HIST", hist), ("SELLO", sello)):
        html, n = re.subn(r"/\*%s_INI\*/[\s\S]*?/\*%s_FIN\*/" % (marca, marca),
                          lambda _m, d=dato, k=marca: "/*%s_INI*/%s/*%s_FIN*/" % (k, seguro(d), k),
                          html, count=1)
        if n != 1:
            sys.exit("La plantilla no tiene los marcadores %s_INI / %s_FIN." % (marca, marca))

    with open(SALIDA, "w", encoding="utf-8") as f:
        f.write(html)

    print("Fuente: %s" % os.path.basename(ruta))
    print("Plantillas: %d · entregadas a Oracle: %d · en proceso: %d · sin iniciar: %d"
          % (len(filas), hist[-1]["ent"], hist[-1]["proceso"], hist[-1]["sin"]))
    print("Corte: %s" % sello["fecha"])


if __name__ == "__main__":
    main()
